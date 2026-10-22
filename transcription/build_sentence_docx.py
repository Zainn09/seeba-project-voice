#!/usr/bin/env python3
"""
Build the sentence-by-sentence Word document + the Roman-Urdu word list.

Outputs (Roman Urdu only):
  all_voices_roman_urdu_sentences.docx   Word document, one numbered
                                         sentence per paragraph
  all_voices_roman_urdu_sentences.txt    plain-text mirror of the docx
  roman_urdu_words.txt                   unique Roman Urdu words,
                                         alphabetical, "word — count"

Sources (nothing is lost or silently merged):
  Part 1 — the 164 human entries of text_voices.docx (Voices 1-162, 166,
           167), split into sentences on . ? !  (Voice 13 has no Roman
           Urdu in the original, so its Urdu sentences are transliterated
           here and clearly marked as generated.)
  Part 2 — the 350 audio files Voice 151-500 transcribed from the audio
           itself (Whisper large-v3). ASR Urdu contains almost no
           sentence punctuation, so each Whisper utterance segment
           (median 8 words) becomes one sentence; terminal punctuation
           is normalised to a full stop where missing.

The document numbers 151-162, 166, 167 therefore appear in BOTH parts
with different content — that is intentional and labelled, because the
.docx entry numbers do not match the same-numbered audio files (see
TRANSCRIPTION_README.md).

Requires: python-docx  (pip install python-docx)
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_output import load_results  # noqa: E402
from transliterate import Transliterator  # noqa: E402
from transcribe import discover  # noqa: E402

SPLIT_RE = re.compile(r"(?<=[.!?])\s+")
WORD_RE = re.compile(r"[a-z]+(?:['-][a-z]+)*")


def split_sentences(text: str) -> list[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    return [s.strip() for s in SPLIT_RE.split(text) if s.strip()]


def ensure_terminal(s: str) -> str:
    s = s.strip()
    return s if s.endswith((".", "?", "!")) else s + "."


def assemble(existing: dict, results: dict, files: dict, tr: Transliterator):
    """Return (part1_entries, part2_entries); entry = dict(heading, note, sentences)."""
    part1 = []
    for n in sorted(int(k) for k in existing):
        e = existing[str(n)]
        heading = e.get("heading") or f"Voice {n}"
        if e.get("roman"):
            sentences = [ensure_terminal(s) for s in split_sentences(e["roman"])]
            note = ""
        elif e.get("urdu"):
            # Voice 13: Roman paragraph blank in the original document.
            urdu_sents = [s for s in re.split(r"(?<=[۔؟!])\s+", e["urdu"]) if s.strip()]
            sentences = []
            for u in urdu_sents:
                roman, _ = tr.transliterate(u)
                sentences += [ensure_terminal(s) for s in split_sentences(roman)]
            note = "Roman Urdu generated from the Urdu text (blank in the original document)."
        else:
            sentences, note = ["[unclear]"], "No text in the original document."
        part1.append({"heading": heading, "note": note, "sentences": sentences})

    part2 = []
    for n in sorted(k for k in files if k >= 151):
        rec = results.get(n)
        heading = f"Voice {n} — {files[n].name}"
        sentences: list[str] = []
        if rec and rec.get("status") == "ok" and rec.get("urdu_text"):
            segments = rec.get("segments") or []
            if segments:
                for seg in segments:
                    roman, _ = tr.transliterate(seg["text"])
                    sentences += [ensure_terminal(s) for s in split_sentences(roman)]
            else:
                roman, _ = tr.transliterate(rec["urdu_text"])
                sentences = [ensure_terminal(s) for s in split_sentences(roman)]
        if not sentences:
            sentences = ["[unclear]"]
        part2.append({"heading": heading, "note": "", "sentences": sentences})
    return part1, part2


def add_page_number_footer(document) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    footer = document.sections[0].footer
    para = footer.paragraphs[0]
    para.alignment = 2  # right
    run = para.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.append(fld)


def build_docx(path: Path, part1, part2, stats: dict) -> None:
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading("All Voices — Roman Urdu Transcriptions", level=0)
    sub = doc.add_paragraph()
    run = sub.add_run("Sentence by sentence  •  Roman Urdu only")
    run.bold = True
    run.font.size = Pt(13)

    for line in [
        f"Generated: {stats['generated']}  |  Entries: {stats['entries']}  |  "
        f"Sentences: {stats['sentences']:,}  |  Unique Roman Urdu words: {stats['unique_words']:,}",
        "Part 1 contains the original human transcriptions from text_voices.docx "
        "(Voices 1-162, 166, 167). Part 2 contains the audio files Voice 151-500, "
        "transcribed from the audio itself with Whisper large-v3 and converted to "
        "Roman Urdu. Voice numbers 151-162, 166 and 167 appear in both parts with "
        "different content, because the entry numbers in text_voices.docx do not "
        "match the same-numbered audio files — both versions are kept, clearly "
        "labelled, and neither overwrites the other.",
        "Voice 126 exists only in Part 1 (its audio file is missing from the "
        "repository). Voice 13 has no Roman Urdu in the original document; its "
        "Roman sentences here are generated from its Urdu text and marked as such. "
        "In Part 2 each numbered item is one spoken sentence/utterance; many "
        "45-second audio slices begin or end mid-sentence, exactly as recorded.",
        "Companion files: roman_urdu_words.txt (unique words with counts), "
        "roman_urdu_transcriptions_151_onward.txt (paragraph format), "
        "transcription/qc_report.txt (quality notes). Original audio files and "
        "text_voices.docx are unchanged.",
    ]:
        doc.add_paragraph(line)

    for part_title, entries in [
        ("Part 1 — From text_voices.docx (original human transcriptions)", part1),
        ("Part 2 — From audio files Voice 151–500 (Whisper large-v3)", part2),
    ]:
        doc.add_heading(part_title, level=1)
        for entry in entries:
            doc.add_heading(entry["heading"], level=2)
            if entry["note"]:
                p = doc.add_paragraph()
                r = p.add_run(entry["note"])
                r.italic = True
                r.font.size = Pt(9)
            for i, sentence in enumerate(entry["sentences"], 1):
                p = doc.add_paragraph(f"{i}. {sentence}")
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(2)

    doc.core_properties.title = "All Voices — Roman Urdu Transcriptions (Sentence by Sentence)"
    doc.core_properties.author = "seeba-project-voice transcription pipeline"
    doc.core_properties.comments = "Part 1: text_voices.docx human transcriptions. Part 2: Whisper large-v3 ASR of audio 151-500."
    try:
        add_page_number_footer(doc)
    except Exception:
        pass  # footer is cosmetic only
    doc.save(str(path))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--voices-dir", default="Voices")
    ap.add_argument("--results-dir", default="transcription/results")
    ap.add_argument("--existing-json", default="transcription/existing_transcriptions.json")
    ap.add_argument("--dictionary", default="transcription/roman_dictionary.json")
    ap.add_argument("--out-docx", default="all_voices_roman_urdu_sentences.docx")
    ap.add_argument("--out-txt", default="all_voices_roman_urdu_sentences.txt")
    ap.add_argument("--out-words", default="roman_urdu_words.txt")
    args = ap.parse_args()

    existing = json.loads(Path(args.existing_json).read_text(encoding="utf-8"))
    results = load_results(Path(args.results_dir))
    files = dict(discover(Path(args.voices_dir)))
    tr = Transliterator.load(Path(args.dictionary))

    part1, part2 = assemble(existing, results, files, tr)
    all_sentences = [s for e in part1 + part2 for s in e["sentences"]]

    counts: Counter = Counter()
    for s in all_sentences:
        if s == "[unclear]":
            continue
        counts.update(WORD_RE.findall(s.lower()))

    stats = {
        "generated": _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%d"),
        "entries": len(part1) + len(part2),
        "sentences": len(all_sentences),
        "unique_words": len(counts),
        "tokens": sum(counts.values()),
    }

    build_docx(Path(args.out_docx), part1, part2, stats)

    # Plain-text mirror (identical structure, easy to diff/preview).
    lines = [
        "All Voices — Roman Urdu Transcriptions (Sentence by Sentence)",
        f"Entries: {stats['entries']} | Sentences: {stats['sentences']} | "
        f"Unique words: {stats['unique_words']} | Total words: {stats['tokens']}",
        "",
    ]
    for part_title, entries in [
        ("Part 1 — From text_voices.docx (original human transcriptions)", part1),
        ("Part 2 — From audio files Voice 151-500 (Whisper large-v3)", part2),
    ]:
        lines += [part_title, "=" * len(part_title), ""]
        for entry in entries:
            lines.append(entry["heading"])
            if entry["note"]:
                lines.append(f"({entry['note']})")
            lines += [f"{i}. {s}" for i, s in enumerate(entry["sentences"], 1)]
            lines.append("")
    Path(args.out_txt).write_text("\n".join(lines), encoding="utf-8")

    # Word list: only words (one per line), alphabetical, with counts.
    words_text = "".join(f"{w} — {counts[w]}\n" for w in sorted(counts))
    Path(args.out_words).write_text(words_text, encoding="utf-8")

    print(f"Part 1 (docx):  {len(part1)} voices, {sum(len(e['sentences']) for e in part1)} sentences")
    print(f"Part 2 (audio): {len(part2)} voices, {sum(len(e['sentences']) for e in part2)} sentences")
    print(f"Total: {stats['entries']} entries, {stats['sentences']:,} sentences, "
          f"{stats['tokens']:,} words, {stats['unique_words']:,} unique words")
    print(f"Top 15 words: {counts.most_common(15)}")
    print(f"Wrote {args.out_docx}\nWrote {args.out_txt}\nWrote {args.out_words}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
